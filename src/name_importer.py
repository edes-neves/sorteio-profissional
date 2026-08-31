"""Names importer — parses names from .csv, .txt, .doc, .docx, .pdf and .odt files.

The parsing libraries are imported lazily so the rest of the application
keeps working even when an optional dependency is not installed. Missing
libraries raise NameImportError with a clear message instead of crashing.
"""

import csv
import os
import re
import subprocess
import tempfile
import zipfile

SUPPORTED_EXTS = (".csv", ".txt", ".doc", ".docx", ".pdf", ".odt")


class NameImportError(Exception):
    pass


def _clean(token: str) -> str:
    """Cleans a single token into a displayable name."""
    token = token.strip()
    token = re.sub(r"[\u00a0\u200b\ufeff]+", " ", token)
    token = re.sub(r"\s+", " ", token)
    return token.strip("\"'`*_-|.,;: \t")


def _extract_names(text: str) -> list[str]:
    """Turns raw extracted text into a list of unique non-empty names.

    Splits on line breaks, commas, semicolons, pipes and tabs so that both
    one-name-per-line and comma-separated files are handled.
    """
    text = text.replace("\u00a0", " ").replace("\ufeff", "")
    if not text.strip():
        return []

    parts = re.split(r"[\n\r\t;|]+", text)
    names: list[str] = []
    for part in parts:
        part = part.strip()
        if not part:
            continue
        for piece in part.split(","):
            name = _clean(piece)
            if name and name not in names:
                names.append(name)
    return names


def _dedupe_preserving_order(items: list[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for item in items:
        key = item.casefold()
        if key not in seen:
            seen.add(key)
            result.append(item)
    return result


def _read_csv(filepath: str) -> str:
    with open(filepath, "r", encoding="utf-8-sig", errors="replace") as f:
        reader = csv.reader(f)
        rows = [row for row in reader if row and any(cell.strip() for cell in row)]
        lines = []
        for row in rows:
            lines.append(", ".join(cell for cell in row if cell.strip()))
        return "\n".join(lines)


def _read_txt(filepath: str) -> str:
    with open(filepath, "r", encoding="utf-8-sig", errors="replace") as f:
        return f.read()


def _read_docx(filepath: str) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise NameImportError(
            "Para importar arquivos .docx instale a biblioteca 'python-docx' "
            "(pip install python-docx)."
        ) from e

    doc = Document(filepath)
    lines = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells if c.text and c.text.strip()]
            if cells:
                lines.append(", ".join(cells))
    return "\n".join(lines)


def _read_odt(filepath: str) -> str:
    """Reads an OpenDocument (.odt) file using only the standard library.

    An .odt is a ZIP archive whose main text lives in ``content.xml`` under
    the OpenDocument text namespace.
    """
    try:
        from xml.etree import ElementTree as ET

        with zipfile.ZipFile(filepath) as archive:
            if "content.xml" not in archive.namelist():
                raise NameImportError(
                    "Arquivo .odt inválido: 'content.xml' não encontrado."
                )
            xml = archive.read("content.xml").decode("utf-8", errors="replace")
        root = ET.fromstring(xml)

        textns = "{urn:oasis:names:tc:opendocument:xmlns:text:1.0}"
        lines = []
        for paragraph in root.iter(textns + "p"):
            text = "".join(paragraph.itertext()).strip()
            if text:
                lines.append(text)
        return "\n".join(lines)
    except NameImportError:
        raise
    except Exception as e:
        raise NameImportError(
            f"Não foi possível ler o arquivo .odt: {e}"
        ) from e


def _read_pdf(filepath: str) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except ImportError:
        try:
            from PyPDF2 import PdfReader  # type: ignore
        except ImportError as e:
            raise NameImportError(
                "Para importar arquivos .pdf instale a biblioteca 'pypdf' "
                "(pip install pypdf)."
            ) from e

    reader = PdfReader(filepath)
    pages = []
    for page in reader.pages:
        try:
            text = page.extract_text() or ""
        except Exception:
            text = ""
        if text.strip():
            pages.append(text)
    return "\n".join(pages)


def _maybe_docx(filepath: str) -> bool:
    """Old .doc is a binary OLE format; .docx is a ZIP. This is a quick
    sanity check to route a mislabeled .docx (or .doc saved as docx)."""
    try:
        if zipfile.is_zipfile(filepath):
            with zipfile.ZipFile(filepath) as z:
                return any(
                    name == "word/document.xml"
                    for name in z.namelist()
                )
    except Exception:
        return False
    return False


def _read_doc(filepath: str) -> str:
    """Best-effort reading for the legacy binary Word .doc format."""
    if _maybe_docx(filepath):
        return _read_docx(filepath)

    # 1) textract (cross-format) if available.
    try:
        import textract  # type: ignore
        return textract.process(filepath).decode("utf-8", errors="replace")
    except ImportError:
        pass
    except Exception:
        pass

    # 2) antiword subprocess if installed on the system.
    try:
        if subprocess.run(["antiword", "-q"], capture_output=True).returncode == 0:
            try:
                result = subprocess.run(
                    ["antiword", filepath],
                    capture_output=True,
                )
                if result.returncode == 0:
                    return result.stdout.decode("utf-8", errors="replace")
            except Exception:
                pass
    except OSError:
        pass

    # 3) Try RTF fallback (many legacy .doc files are RTF in disguise).
    try:
        raw = open(filepath, "rb").read().decode("latin-1", errors="replace")
        if "\\rtf" in raw:
            return _rtf_to_text(raw)
    except Exception:
        pass

    raise NameImportError(
        "Não foi possível ler o arquivo .doc. O formato antigo (.doc) requer "
        "o 'textract' ou o conversor 'antiword' instalados no sistema, ou "
        "converta o arquivo para .docx/.pdf e tente novamente."
    )


def _rtf_to_text(raw: str) -> str:
    """A tiny RTF-to-text converter good enough for name lists."""
    text = raw
    text = text.replace("\\par", "\n").replace("\\line", "\n")
    text = text.replace("\\tab", " ")
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", "", text)
    text = re.sub(r"[{}]", "", text)
    text = re.sub(r"[\\\\]+", "", text)
    return text


def import_names(filepath: str) -> list[str]:
    """Reads a name list from a supported file and returns unique names."""
    if not os.path.isfile(filepath):
        raise NameImportError(f"Arquivo não encontrado: {filepath}")

    ext = os.path.splitext(filepath)[1].lower()
    if ext not in SUPPORTED_EXTS:
        raise NameImportError(
            f"Formato não suportado: '{ext}'. "
            f"Use um dos: {', '.join(SUPPORTED_EXTS)}"
        )

    if ext == ".csv":
        raw_text = _read_csv(filepath)
    elif ext == ".txt":
        raw_text = _read_txt(filepath)
    elif ext == ".docx":
        raw_text = _read_docx(filepath)
    elif ext == ".pdf":
        raw_text = _read_pdf(filepath)
    elif ext == ".odt":
        raw_text = _read_odt(filepath)
    elif ext == ".doc":
        raw_text = _read_doc(filepath)
    else:
        raise NameImportError("Formato de arquivo não suportado.")

    names = _extract_names(raw_text)
    names = _dedupe_preserving_order(names)

    if not names:
        raise NameImportError(
            "Nenhum nome encontrado no arquivo. Verifique se o conteúdo "
            "contém nomes separados por linha ou vírgula."
        )

    return names
